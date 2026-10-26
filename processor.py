import asyncio
import json
import logging
import os
from pathlib import Path
from typing import Callable, Optional

import httpx
import openpyxl
from dotenv import load_dotenv

load_dotenv()

logger = logging.getLogger(__name__)

API_BASE = os.getenv("API_BASE_URL", "https://api.myorg.uz")
API_TOKEN = os.getenv("API_TOKEN")
RATE_LIMIT = 3  # requests per second

# Base columns — founders are appended dynamically after fetching
COLUMNS = [
    "inn", "id", "uuid", "name", "short_name",
    "oked_code", "oked_name",
    "opf_code", "opf_name",
    "activity_state",
    "activity_state_detail_name", "activity_state_detail_name_uz", "activity_state_detail_name_en",
    "business_type_detail_name", "business_type_detail_name_uz", "business_type_detail_name_en",
    "director", "director_uuid",
    "email", "email_status",
    "phones",
    "address",
    "registration_authority", "registration_date", "registration_number",
    "relevance_date", "actual_date",
    "small_businesses",
    "soato_code", "soato_name",
    "soogu_code", "soogu_name",
    "statutory_fund",
    "tax_mode", "trust", "score",
    "vat_number",
    "village_code", "village_name", "village_detail_name",
    "is_bankrupt", "is_abuse_vat", "itpark", "is_suppliers",
    "liquidation_date", "liquidation_reason",
    "connections_all", "connections_director", "connections_founders", "connections_entrepreneur",
    "courts_total", "court_relevance_date",
    "deals_customer_total", "deals_customer_rows",
    "deals_provider_total", "deals_provider_rows",
    "deals_actual_date", "deal_relevance_date",
    "licenses_total", "licenses_relevance_date", "licenses_actual_date",
    "kfs_code", "kfs_name", "kfs_name_ru", "kfs_name_uz_cyr", "kfs_name_uz_lat",
    "billing_region_name", "billing_district_name", "billing_street_name",
    "billing_house", "billing_flat", "billing_postcode",
    "dishonest_executor", "dishonest_executor_delete_date",
    "buildings_total",
    "cadastres_total", "cadastres_relevance_date",
]

NUMBER_COLS = {
    "id", "activity_state", "email_status", "tax_mode", "score",
    "village_code", "is_bankrupt", "is_abuse_vat", "itpark", "is_suppliers",
    "connections_all", "connections_director", "connections_founders",
    "connections_entrepreneur", "courts_total", "deals_customer_total",
    "deals_provider_total", "licenses_total", "buildings_total",
    "cadastres_total", "dishonest_executor", "kfs_code",
}

DATE_COLS = {
    "registration_date", "relevance_date", "actual_date",
    "court_relevance_date", "deal_relevance_date", "liquidation_date",
    "licenses_relevance_date", "licenses_actual_date",
    "cadastres_relevance_date", "deals_actual_date",
    "dishonest_executor_delete_date",
}


def flatten_record(data: dict) -> tuple[dict, list]:
    """Returns (flat_row_dict, founders_list)."""
    row = {}

    for f in [
        "inn", "id", "uuid", "name", "short_name",
        "oked_code", "oked_name", "opf_code", "opf_name",
        "activity_state", "address", "director", "director_uuid",
        "email", "email_status", "registration_authority",
        "registration_date", "registration_number",
        "relevance_date", "actual_date", "small_businesses",
        "soato_code", "soato_name", "soogu_code", "soogu_name",
        "statutory_fund", "tax_mode", "trust", "score",
        "vat_number", "village_code", "village_name",
        "is_bankrupt", "is_abuse_vat", "itpark", "is_suppliers",
        "liquidation_date", "liquidation_reason",
        "court_relevance_date", "deal_relevance_date",
    ]:
        row[f] = data.get(f)

    phones = data.get("phones")
    row["phones"] = ", ".join(str(p) for p in phones) if phones else None

    conn = data.get("connections") or {}
    row["connections_all"] = conn.get("all")
    row["connections_director"] = conn.get("director")
    row["connections_founders"] = conn.get("founders")
    row["connections_entrepreneur"] = conn.get("entrepreneur")

    courts = data.get("courts") or {}
    row["courts_total"] = courts.get("total")

    deals = data.get("deals") or {}
    customer = deals.get("customer") or {}
    provider = deals.get("provider") or {}
    row["deals_customer_total"] = customer.get("total")
    row["deals_customer_rows"] = (
        json.dumps(customer.get("rows"), ensure_ascii=False) if customer.get("rows") else None
    )
    row["deals_provider_total"] = provider.get("total")
    row["deals_provider_rows"] = (
        json.dumps(provider.get("rows"), ensure_ascii=False) if provider.get("rows") else None
    )
    row["deals_actual_date"] = deals.get("actual_date")

    lic = data.get("licenses") or {}
    row["licenses_total"] = lic.get("total")
    row["licenses_relevance_date"] = lic.get("relevance_date")
    row["licenses_actual_date"] = lic.get("actual_date")

    kfs = data.get("kfs") or {}
    row["kfs_code"] = kfs.get("code")
    row["kfs_name"] = kfs.get("name")
    row["kfs_name_ru"] = kfs.get("name_ru")
    row["kfs_name_uz_cyr"] = kfs.get("name_uz_cyr")
    row["kfs_name_uz_lat"] = kfs.get("name_uz_lat")

    asd = data.get("activity_state_detail") or {}
    row["activity_state_detail_name"] = asd.get("name")
    row["activity_state_detail_name_en"] = asd.get("name_en")
    row["activity_state_detail_name_uz"] = asd.get("name_uz")

    btd = data.get("business_type_detail") or {}
    row["business_type_detail_name"] = btd.get("name")
    row["business_type_detail_name_en"] = btd.get("name_en")
    row["business_type_detail_name_uz"] = btd.get("name_uz")

    cba = data.get("company_billing_address") or {}
    row["billing_region_name"] = cba.get("region_name")
    row["billing_district_name"] = cba.get("district_name")
    row["billing_street_name"] = cba.get("street_name")
    row["billing_house"] = cba.get("house")
    row["billing_flat"] = cba.get("flat")
    row["billing_postcode"] = cba.get("postcode")

    vd = data.get("village_detail") or {}
    row["village_detail_name"] = vd.get("name")

    de = data.get("dishonest_executor") or {}
    row["dishonest_executor"] = de.get("is_dishonest_executor")
    row["dishonest_executor_delete_date"] = de.get("delete_date")

    buildings = data.get("buildings") or {}
    row["buildings_total"] = buildings.get("total")

    cadastres = data.get("cadastres") or {}
    row["cadastres_total"] = cadastres.get("total")
    row["cadastres_relevance_date"] = (
        str(cadastres["relevance_date"]) if cadastres.get("relevance_date") else None
    )

    founders = data.get("founders") or []
    return row, founders


def _trademark_columns(n: int) -> list[str]:
    return [f"trademark_{i}_transliteration" for i in range(1, n + 1)]


def _connection_columns(n: int) -> list[str]:
    cols = []
    for i in range(1, n + 1):
        cols += [
            f"connection_{i}_inn",
            f"connection_{i}_name",
            f"connection_{i}_director",
        ]
    return cols


def _founder_columns(n: int) -> list[str]:
    cols = []
    for i in range(1, n + 1):
        cols += [
            f"founder_{i}_name",
            f"founder_{i}_percentage",
            f"founder_{i}_type",
            f"founder_{i}_pinfl",
        ]
    return cols


def detect_inn_column(filepath: str) -> Optional[int]:
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
    wb.close()
    for i, h in enumerate(headers, start=1):
        if h and str(h).strip().upper() in ("INN", "ИНН", "INN_NUMBER", "INNUMBER"):
            return i
    return None


def find_column_by_name(filepath: str, name: str) -> Optional[int]:
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
    wb.close()
    for i, h in enumerate(headers, start=1):
        if h and str(h).strip().lower() == name.strip().lower():
            return i
    return None


def column_letter_to_index(letter: str) -> int:
    result = 0
    for ch in letter.upper():
        result = result * 26 + (ord(ch) - ord("A") + 1)
    return result


def _normalize_inn(val) -> Optional[str]:
    """Excel stores numeric INNs as floats (e.g. 308839903.0) — strip the trailing .0."""
    if val is None:
        return None
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    s = str(val).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s or None


def read_inns(filepath: str, col_idx: int) -> list[str]:
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    inns = []
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, values_only=True):
        inn = _normalize_inn(row[0])
        if inn:
            inns.append(inn)
    wb.close()
    return inns


class RateLimiter:
    """Spaces out request starts to RATE_LIMIT per second, regardless of response time."""
    def __init__(self, rate: float):
        self._interval = 1.0 / rate
        self._lock = asyncio.Lock()
        self._last = 0.0

    async def acquire(self):
        async with self._lock:
            now = asyncio.get_event_loop().time()
            wait = self._last + self._interval - now
            if wait > 0:
                await asyncio.sleep(wait)
            self._last = asyncio.get_event_loop().time()


async def fetch_inn(client: httpx.AsyncClient, inn: str, limiter: RateLimiter) -> tuple[str, dict]:
    await limiter.acquire()
    try:
        resp = await client.get(
            f"/api/entity/inn/{inn}",
            headers={"token": f"Bearer {API_TOKEN}"},
            timeout=30.0,
        )
        if resp.status_code == 200:
            body = resp.json()
            logger.debug("INN %s — rows=%d", inn, len(body.get("rows") or []))
            return inn, body
        logger.error("INN %s — HTTP %s: %s", inn, resp.status_code, resp.text[:300])
    except Exception as e:
        logger.error("INN %s — request failed: %s", inn, e)
    return inn, {}



async def fetch_trademarks(client: httpx.AsyncClient, entity_id, limiter: RateLimiter) -> list:
    if not entity_id:
        return []
    await limiter.acquire()
    try:
        resp = await client.get(
            f"/api/trademark/{entity_id}",
            headers={"token": f"Bearer {API_TOKEN}"},
            timeout=30.0,
        )
        if resp.status_code == 200:
            return resp.json().get("rows") or []
        logger.error("entity %s trademarks — HTTP %s: %s", entity_id, resp.status_code, resp.text[:300])
    except Exception as e:
        logger.error("entity %s trademarks — request failed: %s", entity_id, e)
    return []


async def fetch_connections(client: httpx.AsyncClient, inn: str, limiter: RateLimiter) -> list:
    await limiter.acquire()
    try:
        resp = await client.get(
            f"/api/bank/clients/{inn}/connections",
            headers={"token": f"Bearer {API_TOKEN}"},
            timeout=30.0,
        )
        if resp.status_code == 200:
            return resp.json().get("rows") or []
        logger.error("INN %s connections — HTTP %s: %s", inn, resp.status_code, resp.text[:300])
    except Exception as e:
        logger.error("INN %s connections — request failed: %s", inn, e)
    return []


async def process_excel(
    input_path: str,
    output_path: str,
    col_idx: int,
    progress_callback: Optional[Callable] = None,
) -> int:
    inns = read_inns(input_path, col_idx)
    return await process_inns(inns, output_path, progress_callback)


async def process_inns(
    inns: list[str],
    output_path: str,
    progress_callback: Optional[Callable] = None,
) -> int:
    total = len(inns)

    if total == 0:
        raise ValueError("No INN numbers provided.")

    limiter = RateLimiter(RATE_LIMIT)
    completed = [0]
    raw_results: list[tuple[str, dict, list, list, list]] = [None] * total

    async def fetch_and_store(idx: int, inn: str, client: httpx.AsyncClient):
        try:
            _, data = await fetch_inn(client, inn, limiter)
            flat, founders = flatten_record(data) if data else ({}, [])
            conn_rows = await fetch_connections(client, inn, limiter)
            conn_founders = (conn_rows[0].get("founders") or []) if conn_rows else []
            if conn_founders:
                founders = conn_founders
            trademarks = await fetch_trademarks(client, flat.get("id"), limiter)
            raw_results[idx] = (inn, flat, founders, trademarks, conn_rows)
        except Exception as e:
            logger.error("Unexpected error for INN %s: %s", inn, e)
            raw_results[idx] = (inn, {}, [], [], [])
        completed[0] += 1
        if progress_callback:
            await progress_callback(completed[0], total)

    async with httpx.AsyncClient(base_url=API_BASE) as client:
        await asyncio.gather(*[
            fetch_and_store(i, inn, client) for i, inn in enumerate(inns)
        ], return_exceptions=True)

    max_founders = max((len(founders) for _, _, founders, _, _ in raw_results), default=0)
    max_trademarks = max((len(trademarks) for _, _, _, trademarks, _ in raw_results), default=0)
    max_connections = max((len(conn_rows) for _, _, _, _, conn_rows in raw_results), default=0)
    all_columns = COLUMNS + _founder_columns(max_founders) + _trademark_columns(max_trademarks) + _connection_columns(max_connections)

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.append(all_columns)

    for inn, flat, founders, trademarks, conn_rows in raw_results:
        if flat:
            row_values = [flat.get(c) for c in COLUMNS]
            for f in founders:
                row_values += [
                    f.get("name"),
                    f.get("percentage"),
                    f.get("is_individual"),
                    f.get("pinfl"),
                ]
            row_values += [""] * ((max_founders - len(founders)) * 4)
            for tm in trademarks:
                row_values += [tm.get("transliteration")]
            row_values += [""] * (max_trademarks - len(trademarks))
            for conn in conn_rows:
                row_values += [
                    conn.get("inn"),
                    conn.get("name"),
                    conn.get("director"),
                ]
            row_values += [""] * ((max_connections - len(conn_rows)) * 3)
        else:
            row_values = [""] * len(all_columns)
            row_values[0] = inn  # inn is first column

        ws_out.append([v if (v is not None and v != "") else "—" for v in row_values])

    _format_sheet(ws_out, all_columns)

    wb_out.save(output_path)
    return total


def _format_sheet(ws, columns: list[str]):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    HEADER_COLOR = "1F4E79"
    HEADER_FONT_COLOR = "FFFFFF"

    header_fill = PatternFill("solid", fgColor=HEADER_COLOR)
    header_font = Font(bold=True, color=HEADER_FONT_COLOR, size=11)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    col_widths = {}

    for row_idx, row in enumerate(ws.iter_rows(), start=1):
        for cell in row:
            col_letter = get_column_letter(cell.column)
            col_name = columns[cell.column - 1] if cell.column <= len(columns) else ""

            if row_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell_len = len(str(cell.value)) + 4
            else:
                cell.alignment = left
                if col_name in NUMBER_COLS and cell.value is not None:
                    cell.number_format = "#,##0"
                elif col_name in DATE_COLS and cell.value is not None:
                    cell.number_format = "YYYY-MM-DD"
                cell_len = len(str(cell.value)) if cell.value is not None else 0

            col_widths[col_letter] = max(col_widths.get(col_letter, 0), cell_len)

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = min(width + 2, 60)

    ws.freeze_panes = "A2"
